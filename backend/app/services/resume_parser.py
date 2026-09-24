from pathlib import Path
from typing import Tuple
import pypdf
import docx

class ResumeParser:
    @staticmethod
    def extract_text(file_path: Path) -> Tuple[str, str]:
        """
        Extract text from PDF, DOCX, or TXT/MD files.
        Returns: (extracted_text, file_type)
        """
        suffix = file_path.suffix.lower()
        if suffix == ".pdf":
            return ResumeParser._parse_pdf(file_path), "PDF"
        elif suffix in [".docx", ".doc"]:
            return ResumeParser._parse_docx(file_path), "DOCX"
        elif suffix in [".txt", ".md"]:
            return ResumeParser._parse_text(file_path), "TEXT"
        else:
            raise ValueError(f"Unsupported file format: {suffix}")

    @staticmethod
    def _parse_pdf(file_path: Path) -> str:
        text_parts = []
        with open(file_path, "rb") as f:
            reader = pypdf.PdfReader(f)
            for page_num, page in enumerate(reader.pages):
                text = page.extract_text()
                if text:
                    text_parts.append(text.strip())
        return "\n\n".join(text_parts)

    @staticmethod
    def _parse_docx(file_path: Path) -> str:
        doc = docx.Document(file_path)
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also parse tables if any
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    paragraphs.append(row_text)
        return "\n".join(paragraphs)

    @staticmethod
    def _parse_text(file_path: Path) -> str:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()
